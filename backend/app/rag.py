"""Ingestion, chunking, embedding and hybrid retrieval."""
import csv
import io
import os
import re
from dataclasses import dataclass

from fastembed import TextEmbedding
from pgvector.psycopg import register_vector
from psycopg.rows import dict_row
from psycopg_pool import ConnectionPool

from app import heavy

EMBED_MODEL = os.getenv("EMBED_MODEL", "BAAI/bge-small-en-v1.5")  # 384 dims, matches schema.sql
CHUNK_SIZE, OVERLAP = 1100, 150
MIN_PAGE_CHARS = int(os.getenv("MIN_PAGE_CHARS", "25"))      # fewer characters than this = treat the page as a scan
MAX_OCR_PAGES = int(os.getenv("MAX_OCR_PAGES", "30"))        # OCR is slow on a free CPU, so cap it per file
RERANK_CANDIDATES = 20

_embedder = None
pool = None


class Unsupported(Exception):
    pass


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None


def init():
    global _embedder, pool
    _embedder = TextEmbedding(EMBED_MODEL, cache_dir=os.getenv("FASTEMBED_CACHE_PATH"))
    list(_embedder.embed(["warm up"]))
    pool = ConnectionPool(
        os.environ["DATABASE_URL"],
        min_size=1,
        max_size=4,
        kwargs={"autocommit": True, "prepare_threshold": None, "row_factory": dict_row},
        configure=register_vector,
        open=True,
    )


# ---------- parsing ----------
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
              ".gif": "image/gif", ".webp": "image/webp"}


SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".tsv", ".html", ".htm", ".txt", ".md",
             ".json", ".xml", ".log", ".py", ".js", ".ts", ".yaml", ".yml", ".rtf"} | set(IMAGE_MIME)


def is_supported(filename):
    return os.path.splitext(filename)[1].lower() in SUPPORTED


def _text(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def _rows_to_blocks(rows, section):
    rows = iter(rows)
    header = next(rows, None)
    if not header:
        return []
    names = [str(h).strip() if h not in (None, "") else f"col{i+1}" for i, h in enumerate(header)]
    lines, out = [], []
    for row in rows:
        cells = [f"{n}: {v}" for n, v in zip(names, row) if v not in (None, "")]
        if cells:
            lines.append("; ".join(cells))
        if len(lines) == 15:
            out.append(Block("\n".join(lines), None, section))
            lines = []
    if lines:
        out.append(Block("\n".join(lines), None, section))
    return out


def parse(filename, data, describe_image, progress=lambda m: None, notes=None):
    notes = notes if notes is not None else []
    ext = os.path.splitext(filename)[1].lower()
    blocks = []
    if ext == ".pdf":
        from pypdf import PdfReader
        texts = [(p.extract_text() or "").strip() for p in PdfReader(io.BytesIO(data)).pages]
        scanned = [i + 1 for i, t in enumerate(texts) if len(t) < MIN_PAGE_CHARS]
        ocr = {}
        if scanned and heavy.OCR_ON:
            todo = scanned[:MAX_OCR_PAGES]
            ocr = heavy.ocr_pdf_pages(data, todo, lambda k, n: progress(f"Reading scanned page {k + 1} of {n}…"))
            if len(scanned) > len(todo):
                notes.append(f"Read {len(todo)} of {len(scanned)} scanned pages (limit).")
        for i, t in enumerate(texts, 1):
            if i in ocr and len(ocr[i]) > len(t):
                t = ocr[i]
            if t:
                blocks.append(Block(t, i))
    elif ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        section = None
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            if p.style.name.lower().startswith("heading"):
                section = t
            blocks.append(Block(t, None, section))
        for table in d.tables:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
            blocks.append(Block("\n".join(rows), None, "Table"))
    elif ext == ".pptx":
        from pptx import Presentation
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            parts = [s.text_frame.text.strip() for s in slide.shapes
                     if s.has_text_frame and s.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
            if parts:
                blocks.append(Block("\n".join(parts), i))
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            blocks += _rows_to_blocks(ws.iter_rows(values_only=True), ws.title)
    elif ext in (".csv", ".tsv"):
        delim = "\t" if ext == ".tsv" else ","
        blocks += _rows_to_blocks(csv.reader(io.StringIO(_text(data)), delimiter=delim), "Table")
    elif ext in (".html", ".htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_text(data), "html.parser")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        blocks.append(Block(re.sub(r"\n{3,}", "\n\n", soup.get_text("\n")).strip()))
    elif ext in IMAGE_MIME:
        blocks.append(Block(describe_image(data, IMAGE_MIME[ext]), None, "Image description"))
    elif ext in (".txt", ".md", ".json", ".xml", ".log", ".py", ".js", ".ts", ".yaml", ".yml", ".rtf"):
        blocks.append(Block(_text(data).strip()))
    else:
        raise Unsupported(f"{ext or 'This'} files aren't supported yet.")
    return blocks


# ---------- chunking ----------
def _split(text):
    if len(text) <= CHUNK_SIZE:
        return [text]
    out, start = [], 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            cut = max(text.rfind("\n\n", start, end), text.rfind("\n", start, end), text.rfind(". ", start, end))
            if cut > start + CHUNK_SIZE * 0.5:
                end = cut + 1
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - OVERLAP, start + 1)
    return [c for c in out if c]


def chunk(blocks):
    """Merge small blocks that share a page/section, split large ones with overlap."""
    chunks, buf, meta = [], "", (None, None)

    def flush():
        nonlocal buf
        if buf.strip():
            chunks.extend((c, *meta) for c in _split(buf.strip()))
        buf = ""

    for b in blocks:
        m = (b.page, b.section)
        if m != meta or len(buf) + len(b.text) > CHUNK_SIZE:
            flush()
            meta = m
        buf += ("\n\n" if buf else "") + b.text
    flush()
    return chunks


# ---------- storage ----------
DOC_COLS = "id, filename, n_chunks, status, progress, error, created_at"


def create_doc(ns, filename, mime):
    with pool.connection() as c:
        row = c.execute(
            f"insert into documents(namespace, filename, mime, status, progress) "
            f"values (%s,%s,%s,'processing','Queued…') returning {DOC_COLS}", (ns, filename, mime)).fetchone()
    return {**row, "id": str(row["id"])}


def _progress(doc_id, msg):
    with pool.connection() as c:
        c.execute("update documents set progress=%s where id=%s", (msg, doc_id))


def _fail(doc_id, msg):
    with pool.connection() as c:
        c.execute("update documents set status='failed', error=%s, progress=null where id=%s", (msg, doc_id))


def process(ns, doc_id, filename, data, describe_image):
    """Runs in a background worker: parse (with OCR if needed), chunk, embed, store."""
    try:
        notes = []
        _progress(doc_id, "Extracting text…")
        chunks = chunk(parse(filename, data, describe_image, lambda m: _progress(doc_id, m), notes))
        if not chunks:
            hint = " This looks like a scan and OCR is switched off." if not heavy.OCR_ON else ""
            raise Unsupported("No readable text found." + hint)
        _progress(doc_id, f"Indexing {len(chunks)} parts…")
        # The file name and section go into the embedding so short chunks keep their context.
        texts = [f"{filename} | {sec or ''} | {c}" for c, _, sec in chunks]
        vecs = list(_embedder.passage_embed(texts, batch_size=32))
        with pool.connection() as conn, conn.transaction():
            with conn.cursor() as cur:
                cur.executemany(
                    "insert into chunks(document_id, namespace, ord, page, section, content, embedding) "
                    "values (%s,%s,%s,%s,%s,%s,%s)",
                    [(doc_id, ns, i, page, sec, text, v)
                     for i, ((text, page, sec), v) in enumerate(zip(chunks, vecs))])
            conn.execute("update documents set n_chunks=%s, status='ready', progress=%s, error=null where id=%s",
                         (len(chunks), " ".join(notes) or None, doc_id))
    except Unsupported as e:
        _fail(doc_id, str(e))
    except Exception:
        _fail(doc_id, "That file couldn't be read. It may be corrupted or password protected.")


def reset_stale():
    """Jobs die with the server. Mark anything left unfinished so it doesn't spin forever."""
    with pool.connection() as c:
        c.execute("update documents set status='failed', progress=null, "
                  "error='Interrupted by a restart. Please upload the file again.' where status='processing'")


def count_processing():
    with pool.connection() as c:
        return c.execute("select count(*) as n from documents where status='processing'").fetchone()["n"]


def list_docs(ns):
    with pool.connection() as c:
        rows = c.execute(f"select {DOC_COLS} from documents "
                         "where namespace=%s order by created_at desc", (ns,)).fetchall()
    return [{**r, "id": str(r["id"])} for r in rows]


def count_docs(ns):
    with pool.connection() as c:
        return c.execute("select count(*) as n from documents where namespace=%s", (ns,)).fetchone()["n"]


def delete_doc(ns, doc_id):
    with pool.connection() as c:
        c.execute("delete from documents where id=%s and namespace=%s", (doc_id, ns))


def ping_db():
    with pool.connection() as c:
        c.execute("select 1")


# ---------- hybrid retrieval ----------
def search(ns, query, doc_ids=None, k=30, n=8):
    """Dense (pgvector) + keyword (Postgres full text) fused with reciprocal rank fusion, then reranked."""
    cands = RERANK_CANDIDATES if heavy.RERANK_ON else n
    qv = next(iter(_embedder.query_embed(query)))
    words = list(dict.fromkeys(w.lower() for w in re.findall(r"[A-Za-z0-9]{3,}", query)))[:12]
    flt = "and document_id = any(%(docs)s::uuid[])" if doc_ids else ""
    sql = f"""
    with sem as (
      select id, row_number() over (order by embedding <=> %(qv)s) as r
      from chunks where namespace = %(ns)s {flt}
      order by embedding <=> %(qv)s limit %(k)s),
    lex as (
      select id, row_number() over (order by ts_rank_cd(fts, q) desc) as r
      from chunks, to_tsquery('english', %(tsq)s) q
      where namespace = %(ns)s {flt} and fts @@ q
      order by ts_rank_cd(fts, q) desc limit %(k)s)
    select c.document_id, d.filename, c.page, c.section, c.content,
           coalesce(1.0/(60+sem.r), 0) + coalesce(1.0/(60+lex.r), 0) as score
    from sem full outer join lex on sem.id = lex.id
    join chunks c on c.id = coalesce(sem.id, lex.id)
    join documents d on d.id = c.document_id
    order by score desc limit %(n)s"""
    params = {"qv": qv, "ns": ns, "k": k, "n": cands, "docs": doc_ids,
              "tsq": " | ".join(words) or "xqzxqz"}
    with pool.connection() as c:
        rows = c.execute(sql, params).fetchall()
    hits = [{**r, "document_id": str(r["document_id"]), "score": float(r["score"])} for r in rows]
    return heavy.rerank(query, hits, n)
